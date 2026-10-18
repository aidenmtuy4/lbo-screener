"""Builds Stride_LBO_Model.xlsx -- a formula-driven LBO of Stride, Inc. (NYSE: LRN)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

F = "Arial"
BLUE, BLACK, GREEN, WHITE = "0000FF", "000000", "008000", "FFFFFF"
NAVY = "1F3864"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="808080")

USD = '$#,##0.0;($#,##0.0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x";(0.0"x");"-"'

SRC_10K = "Source: Stride FY2026 earnings release (Aug 2026) / 10-K for FY ended Jun 30, 2026, sec.gov"


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)


def setup(ws, widths):
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(NAVY, True, 14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font("595959", italic=True, size=9)


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = font(WHITE, True)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center" if i else "left")


def section(ws, row, text, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = SUB_FILL
    ws.cell(row=row, column=1, value=text).font = font(NAVY, True)


def put(ws, ref, value, fmt=None, color=BLACK, bold=False, note=None, key=False):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold)
    if fmt:
        c.number_format = fmt
    if note:
        c.comment = Comment(note, "Model")
    if key:
        c.fill = KEY_FILL
    return c


wb = Workbook()

# =============================================================================
# ASSUMPTIONS
# =============================================================================
A = wb.active
A.title = "Assumptions"
setup(A, [44, 14, 70])
title(A, "Stride, Inc. (LRN) -- LBO Assumptions",
      "$ in millions except per share. Blue = hard-coded input, black = formula, yellow = key lever.")
REF = {}
r = 4


def block(name):
    global r
    r += 1
    section(A, r, name, 3)
    r += 1


def inp(key, label, value, fmt, source, key_lever=False):
    global r
    A.cell(row=r, column=1, value=label).font = font()
    put(A, f"B{r}", value, fmt, BLUE, key=key_lever)
    A.cell(row=r, column=3, value=source).font = font("595959", size=9, italic=True)
    REF[key] = f"Assumptions!$B${r}"
    r += 1


header(A, 4, ["Input", "Value", "Source / rationale"])
block("Transaction")
inp("price", "Current share price ($)", 82.39, USD2, "Yahoo Finance chart data, pulled Oct 17, 2026 by market_data.py")
inp("shares", "Shares outstanding (mm)", 41.559845, NUM, "10-K cover page, as of Jul 31, 2026")
inp("premium", "Takeover premium", 0.30, PCT, "Typical US take-private premium range is 25-40%", True)
inp("fees_pct", "Transaction fees (% of EV)", 0.02, PCT, "Advisory + financing fees, standard assumption")
inp("min_cash", "Minimum cash left in business", 100.0, USD, "Operating cash buffer (~4% of revenue)")
block("Balance sheet at close (Jun 30, 2026)")
inp("cash", "Cash and equivalents", 754.501, USD, SRC_10K)
inp("securities", "Marketable securities", 203.499, USD, SRC_10K)
inp("converts", "Convertible senior notes (1.125%, due Sep 2027)", 417.995, USD, SRC_10K + ". Refinanced at close")
inp("leases", "Finance lease liabilities", 116.932, USD, SRC_10K + ". Treated as debt, refinanced at close")
block("FY2026 actuals (fiscal year ended Jun 30, 2026)")
inp("rev_ge", "Revenue -- General Education", 1417.785, USD, SRC_10K)
inp("rev_cl", "Revenue -- Career Learning", 1100.296, USD, SRC_10K)
inp("ebit", "Income from operations (EBIT)", 450.767, USD, SRC_10K)
inp("da", "Depreciation & amortization", 126.562, USD, SRC_10K)
inp("capex", "Capex (PP&E + capitalized software + curriculum)", 78.846, USD, SRC_10K + ": 0.6 + 61.6 + 16.7")
inp("sbc", "Stock-based compensation (memo, NOT added back)", 40.255, USD,
    "Conservative: GAAP EBITDA used, so SBC is treated as a real cost a PE owner pays in cash")
block("Financing")
inp("max_lev", "Max total debt / EBITDA", 5.0, MULT, "Typical leveraged-loan market ceiling for a B-rated borrower", True)
inp("min_eq", "Minimum sponsor equity (% of funding need)", 0.40, PCT, "Lenders typically require 40-50% equity in current market")
inp("cover", "Minimum EBITDA / interest at close", 2.0, MULT, "Rating-agency comfort level")
inp("tlb_share", "Term Loan B share of total debt", 0.70, PCT, "Remainder funded with senior notes")
inp("tlb_rate", "Term Loan B interest rate", 0.0775, PCT, "SOFR (~4.25%) + 350bp spread")
inp("notes_rate", "Senior notes interest rate", 0.09, PCT, "B-rated high-yield context")
inp("amort", "TLB mandatory amortization (% of original / yr)", 0.01, PCT, "Market standard for TLBs")
inp("sweep", "Excess cash flow sweep to TLB", 1.00, PCT, "100% of free cash flow repays the TLB early")
block("Operating & exit")
inp("tax", "Tax rate", 0.25, PCT, "FY2026 effective rate was ~23%; 25% is conservative")
inp("nwc", "Working capital investment (% of revenue change)", 0.10, PCT, "Receivables from school districts grow with revenue")
inp("exit_delta", "Exit multiple change vs. entry (x)", 0.0, MULT, "0 = exit at the entry multiple (no multiple expansion)", True)

# =============================================================================
# LBO MODEL
# =============================================================================
M = wb.create_sheet("LBO")
setup(M, [40, 13, 13, 13, 13, 13, 13, 3, 30])
title(M, "Stride, Inc. (LRN) -- Leveraged Buyout Model", "$ in millions. Fiscal years end June 30.")
YEARS = ["FY2026A", "FY2027E", "FY2028E", "FY2029E", "FY2030E", "FY2031E"]
COLS = ["B", "C", "D", "E", "F", "G"]          # B = actual, C..G = projections
P = COLS[1:]
header(M, 4, [""] + YEARS)


def link(key):
    return "=" + REF[key]


def row_label(rr, text, bold=False, indent=0):
    c = M.cell(row=rr, column=1, value=("   " * indent) + text)
    c.font = font(bold=bold)


# ---- Drivers ----
section(M, 5, "Operating drivers", 7)
GE_G, CL_G, MARGIN, DA_PCT, CX_PCT = 6, 7, 8, 9, 10
row_label(GE_G, "General Education revenue growth", indent=1)
row_label(CL_G, "Career Learning revenue growth", indent=1)
row_label(MARGIN, "EBITDA margin", indent=1)
row_label(DA_PCT, "D&A % of revenue", indent=1)
row_label(CX_PCT, "Capex % of revenue", indent=1)
ge_growth = [-0.01, -0.01, 0.0, 0.0, 0.0]
cl_growth = [0.10, 0.09, 0.08, 0.07, 0.06]
M[f"B{GE_G}"].comment = Comment("FY2026A: Gen Ed revenue -2.1%, Career Learning +15.0% (FY2026 earnings release)", "Model")
put(M, f"B{GE_G}", "=B14/1448.676-1", PCT, note="FY2025 Gen Ed revenue $1,448.7mm (FY2026 earnings release)")
put(M, f"B{CL_G}", "=B15/956.641-1", PCT, note="FY2025 Career Learning revenue $956.6mm (FY2026 earnings release)")
for i, col in enumerate(P):
    put(M, f"{col}{GE_G}", ge_growth[i], PCT, BLUE,
        note="Gen Ed enrollments fell ~2.5% in FY2026; assume modest decline then flat" if i == 0 else None)
    put(M, f"{col}{CL_G}", cl_growth[i], PCT, BLUE,
        note="Career Learning enrollments +14% in FY2026; assume growth decays toward high single digits" if i == 0 else None)
put(M, f"B{MARGIN}", "=B18/B16", PCT)
put(M, f"B{DA_PCT}", "=-B19/B16", PCT)
put(M, f"B{CX_PCT}", "=" + REF["capex"] + "/B16", PCT, GREEN)
for i, col in enumerate(P):
    prev = COLS[i]
    for rr in (MARGIN, DA_PCT, CX_PCT):
        put(M, f"{col}{rr}", f"={prev}{rr}", PCT,
            note="Held flat at FY2026A level -- no margin expansion assumed" if (rr == MARGIN and i == 0) else None)

# ---- Income statement & cash flow ----
section(M, 12, "Income statement & free cash flow", 7)
GE, CL, REV, RG, EBITDA, DA, EBIT, INT, EBT, TAX, NI, ADD_DA, CAPEX, NWC, FCF = range(14, 29)
for rr, lab, b, ind in [(GE, "General Education revenue", 0, 1), (CL, "Career Learning revenue", 0, 1),
                        (REV, "Total revenue", 1, 0), (RG, "Revenue growth", 0, 1),
                        (EBITDA, "EBITDA", 1, 0), (DA, "Less: D&A", 0, 1), (EBIT, "EBIT", 1, 0),
                        (INT, "Less: Interest expense", 0, 1), (EBT, "Pre-tax income", 0, 0),
                        (TAX, "Less: Taxes", 0, 1), (NI, "Net income", 1, 0),
                        (ADD_DA, "Plus: D&A", 0, 1), (CAPEX, "Less: Capex", 0, 1),
                        (NWC, "Less: Increase in working capital", 0, 1), (FCF, "Levered free cash flow", 1, 0)]:
    row_label(rr, lab, bool(b), ind)

put(M, f"B{GE}", link("rev_ge"), USD, GREEN)
put(M, f"B{CL}", link("rev_cl"), USD, GREEN)
put(M, f"B{REV}", f"=B{GE}+B{CL}", USD, bold=True)
put(M, f"B{RG}", f"=B{REV}/2405.317-1", PCT, note="FY2025 total revenue $2,405.3mm (FY2026 earnings release)")
put(M, f"B{EBITDA}", "=" + REF["ebit"] + "+" + REF["da"], USD, GREEN, True)
put(M, f"B{DA}", "=-" + REF["da"], USD, GREEN)
put(M, f"B{EBIT}", f"=B{EBITDA}+B{DA}", USD, bold=True)

TLB_BEG, TLB_MAND, TLB_SWEEP, TLB_END, NOTES, TOT_DEBT, CASH, NET_DEBT, LEV, NLEV, COVER = range(31, 42)
for i, col in enumerate(P):
    prev = COLS[i]
    put(M, f"{col}{GE}", f"={prev}{GE}*(1+{col}{GE_G})", USD)
    put(M, f"{col}{CL}", f"={prev}{CL}*(1+{col}{CL_G})", USD)
    put(M, f"{col}{REV}", f"={col}{GE}+{col}{CL}", USD, bold=True)
    put(M, f"{col}{RG}", f"={col}{REV}/{prev}{REV}-1", PCT)
    put(M, f"{col}{EBITDA}", f"={col}{REV}*{col}{MARGIN}", USD, bold=True)
    put(M, f"{col}{DA}", f"=-{col}{REV}*{col}{DA_PCT}", USD)
    put(M, f"{col}{EBIT}", f"={col}{EBITDA}+{col}{DA}", USD, bold=True)
    put(M, f"{col}{INT}", f"=-({prev}{TLB_END}*{REF['tlb_rate']}+{prev}{NOTES}*{REF['notes_rate']})", USD,
        note="Interest on beginning-of-year balances (avoids a circular reference)" if i == 0 else None)
    put(M, f"{col}{EBT}", f"={col}{EBIT}+{col}{INT}", USD)
    put(M, f"{col}{TAX}", f"=-MAX({col}{EBT},0)*{REF['tax']}", USD)
    put(M, f"{col}{NI}", f"={col}{EBT}+{col}{TAX}", USD, bold=True)
    put(M, f"{col}{ADD_DA}", f"=-{col}{DA}", USD)
    put(M, f"{col}{CAPEX}", f"=-{col}{REV}*{col}{CX_PCT}", USD)
    put(M, f"{col}{NWC}", f"=-({col}{REV}-{prev}{REV})*{REF['nwc']}", USD)
    put(M, f"{col}{FCF}", f"={col}{NI}+{col}{ADD_DA}+{col}{CAPEX}+{col}{NWC}", USD, bold=True)

# ---- Debt schedule ----
section(M, 30, "Debt schedule", 7)
for rr, lab, b, ind in [(TLB_BEG, "Term Loan B -- beginning", 0, 1), (TLB_MAND, "Mandatory amortization", 0, 2),
                        (TLB_SWEEP, "Optional repayment (cash sweep)", 0, 2), (TLB_END, "Term Loan B -- ending", 1, 1),
                        (NOTES, "Senior notes", 0, 1), (TOT_DEBT, "Total debt", 1, 0),
                        (CASH, "Cash balance", 0, 0), (NET_DEBT, "Net debt", 1, 0),
                        (LEV, "Total debt / EBITDA", 0, 1), (NLEV, "Net debt / EBITDA", 0, 1),
                        (COVER, "EBITDA / interest", 0, 1)]:
    row_label(rr, lab, bool(b), ind)

# Sources & Uses rows (defined below) feed the opening balances
SU = 45
OFFER, PURCH, REFI, FEES, USES = SU + 1, SU + 2, SU + 3, SU + 4, SU + 5
CASH_USED, TLB0, NOTES0, EQUITY, SOURCES = SU + 7, SU + 8, SU + 9, SU + 10, SU + 11
ENTRY = SU + 13
ENTRY_EV, LTM, ENTRY_MULT, BLENDED, DEBT0, DEBT_MULT, EQ_PCT = range(ENTRY + 1, ENTRY + 8)

put(M, f"B{TLB_END}", f"=B{TLB0}", USD, bold=True)
put(M, f"B{NOTES}", f"=B{NOTES0}", USD)
put(M, f"B{TOT_DEBT}", f"=B{TLB_END}+B{NOTES}", USD, bold=True)
put(M, f"B{CASH}", link("min_cash"), USD, GREEN)
put(M, f"B{NET_DEBT}", f"=B{TOT_DEBT}-B{CASH}", USD, bold=True)
put(M, f"B{LEV}", f"=B{TOT_DEBT}/B{EBITDA}", MULT)
put(M, f"B{NLEV}", f"=B{NET_DEBT}/B{EBITDA}", MULT)
M[f"A{TLB_BEG}"].comment = Comment("FY2026A column shows pro forma balances at deal close", "Model")
for i, col in enumerate(P):
    prev = COLS[i]
    put(M, f"{col}{TLB_BEG}", f"={prev}{TLB_END}", USD)
    put(M, f"{col}{TLB_MAND}", f"=-MIN($B${TLB0}*{REF['amort']},{col}{TLB_BEG})", USD)
    put(M, f"{col}{TLB_SWEEP}",
        f"=-MIN(MAX({col}{FCF}+{col}{TLB_MAND},0)*{REF['sweep']},{col}{TLB_BEG}+{col}{TLB_MAND})", USD)
    put(M, f"{col}{TLB_END}", f"={col}{TLB_BEG}+{col}{TLB_MAND}+{col}{TLB_SWEEP}", USD, bold=True)
    put(M, f"{col}{NOTES}", f"={prev}{NOTES}", USD)
    put(M, f"{col}{TOT_DEBT}", f"={col}{TLB_END}+{col}{NOTES}", USD, bold=True)
    put(M, f"{col}{CASH}", f"={prev}{CASH}+{col}{FCF}+{col}{TLB_MAND}+{col}{TLB_SWEEP}", USD)
    put(M, f"{col}{NET_DEBT}", f"={col}{TOT_DEBT}-{col}{CASH}", USD, bold=True)
    put(M, f"{col}{LEV}", f"={col}{TOT_DEBT}/{col}{EBITDA}", MULT)
    put(M, f"{col}{NLEV}", f"={col}{NET_DEBT}/{col}{EBITDA}", MULT)
    put(M, f"{col}{COVER}", f"=IF({col}{INT}=0,0,{col}{EBITDA}/-{col}{INT})", MULT)

# ---- Sources & Uses ----
section(M, SU - 1, "Sources & uses at close", 7)
for rr, lab, b, ind in [(OFFER, "Offer price per share ($)", 0, 1), (PURCH, "Purchase of equity", 0, 1),
                        (REFI, "Refinance existing debt", 0, 1), (FEES, "Transaction fees", 0, 1),
                        (USES, "Total uses", 1, 0), (CASH_USED, "Existing cash used", 0, 1),
                        (TLB0, "New Term Loan B", 0, 1), (NOTES0, "New senior notes", 0, 1),
                        (EQUITY, "Sponsor equity", 1, 1), (SOURCES, "Total sources", 1, 0)]:
    row_label(rr, lab, bool(b), ind)
M.cell(row=SU, column=1, value="Uses").font = font(bold=True, italic=True)
M.cell(row=SU + 6, column=1, value="Sources").font = font(bold=True, italic=True)
put(M, f"B{OFFER}", f"={REF['price']}*(1+{REF['premium']})", USD2, GREEN)
put(M, f"B{PURCH}", f"=B{OFFER}*{REF['shares']}", USD, GREEN)
put(M, f"B{REFI}", f"={REF['converts']}+{REF['leases']}", USD, GREEN)
put(M, f"B{FEES}", f"=B{ENTRY_EV}*{REF['fees_pct']}", USD, GREEN)
put(M, f"B{USES}", f"=B{PURCH}+B{REFI}+B{FEES}", USD, bold=True)
put(M, f"B{CASH_USED}", f"=MAX({REF['cash']}+{REF['securities']}-{REF['min_cash']},0)", USD, GREEN)
put(M, f"B{TLB0}", f"=B{DEBT0}*{REF['tlb_share']}", USD, GREEN)
put(M, f"B{NOTES0}", f"=B{DEBT0}-B{TLB0}", USD)
put(M, f"B{EQUITY}", f"=B{USES}-B{CASH_USED}-B{TLB0}-B{NOTES0}", USD, bold=True, key=True)
put(M, f"B{SOURCES}", f"=B{CASH_USED}+B{TLB0}+B{NOTES0}+B{EQUITY}", USD, bold=True)
put(M, f"C{SOURCES}", f'=IF(ABS(B{SOURCES}-B{USES})<0.01,"balances","CHECK")', None)

section(M, ENTRY, "Entry valuation & debt sizing", 7)
for rr, lab in [(ENTRY_EV, "Entry enterprise value"), (LTM, "FY2026A EBITDA"), (ENTRY_MULT, "Entry EV / EBITDA"),
                (BLENDED, "Blended interest rate on new debt"), (DEBT0, "Total new debt"),
                (DEBT_MULT, "New debt / EBITDA"), (EQ_PCT, "Sponsor equity % of funding need")]:
    row_label(rr, lab, rr in (ENTRY_EV, DEBT0), 1)
put(M, f"B{ENTRY_EV}",
    f"=B{PURCH}+{REF['converts']}+{REF['leases']}-{REF['cash']}-{REF['securities']}", USD, GREEN, True,
    note="Equity purchase price + debt - cash - marketable securities")
put(M, f"B{LTM}", f"=B{EBITDA}", USD)
put(M, f"B{ENTRY_MULT}", f"=B{ENTRY_EV}/B{LTM}", MULT, key=True)
put(M, f"B{BLENDED}", f"={REF['tlb_share']}*{REF['tlb_rate']}+(1-{REF['tlb_share']})*{REF['notes_rate']}", PCT, GREEN)
put(M, f"B{DEBT0}",
    f"=MIN({REF['max_lev']}*B{LTM},(1-{REF['min_eq']})*(B{USES}-B{CASH_USED}),B{LTM}/({REF['cover']}*B{BLENDED}))",
    USD, GREEN, True, note="Smallest of: (1) max leverage x EBITDA, (2) amount that keeps sponsor equity at the "
                           "minimum %, (3) amount that keeps EBITDA/interest at the minimum")
put(M, f"B{DEBT_MULT}", f"=B{DEBT0}/B{LTM}", MULT)
put(M, f"B{EQ_PCT}", f"=B{EQUITY}/(B{USES}-B{CASH_USED})", PCT)

# ---- Returns ----
RET = EQ_PCT + 2
section(M, RET, "Returns (exit at end of FY2031)", 7)
EXIT_MULT, EXIT_EBITDA, EXIT_EV, EXIT_ND, EXIT_EQ, MOIC, IRR_ROW = range(RET + 1, RET + 8)
CF = IRR_ROW + 2
for rr, lab, b in [(EXIT_MULT, "Exit EV / EBITDA", 0), (EXIT_EBITDA, "FY2031E EBITDA", 0),
                   (EXIT_EV, "Exit enterprise value", 1), (EXIT_ND, "Less: net debt at exit", 0),
                   (EXIT_EQ, "Equity value at exit", 1), (MOIC, "MOIC", 1), (IRR_ROW, "IRR", 1),
                   (CF, "Sponsor cash flows", 0)]:
    row_label(rr, lab, bool(b), 1)
put(M, f"B{EXIT_MULT}", f"=B{ENTRY_MULT}+{REF['exit_delta']}", MULT, GREEN)
put(M, f"B{EXIT_EBITDA}", f"=G{EBITDA}", USD)
put(M, f"B{EXIT_EV}", f"=B{EXIT_MULT}*B{EXIT_EBITDA}", USD, bold=True)
put(M, f"B{EXIT_ND}", f"=-G{NET_DEBT}", USD)
put(M, f"B{EXIT_EQ}", f"=B{EXIT_EV}+B{EXIT_ND}", USD, bold=True)
put(M, f"B{MOIC}", f"=B{EXIT_EQ}/B{EQUITY}", '0.00"x"', bold=True, key=True)
put(M, f"B{IRR_ROW}", f"=IRR(B{CF}:G{CF})", PCT, bold=True, key=True)
put(M, f"B{CF}", f"=-B{EQUITY}", USD)
for col in P[:-1]:
    put(M, f"{col}{CF}", 0, USD, BLUE, note="No interim dividends assumed" if col == "C" else None)
put(M, f"G{CF}", f"=B{EXIT_EQ}", USD)
for c, y in zip(COLS, YEARS):
    M[f"{c}{CF - 1}"] = y
    M[f"{c}{CF - 1}"].font = font("595959", size=9)
    M[f"{c}{CF - 1}"].alignment = Alignment(horizontal="right")
M.freeze_panes = "B5"

# =============================================================================
# SENSITIVITY  (all formulas: exit multiple x exit year)
# =============================================================================
S = wb.create_sheet("Sensitivity")
setup(S, [30, 13, 13, 13, 3, 30, 13, 13, 13])
title(S, "Returns sensitivity -- exit multiple vs. holding period",
      "Every cell is a live formula off the LBO tab. Debt paydown path doesn't depend on exit multiple.")
deltas = [-2, -1, 0, 1, 2]
exit_years = [("FY2029E", "E", 3), ("FY2030E", "F", 4), ("FY2031E", "G", 5)]
for block_col, label, kind in [(1, "IRR", "irr"), (6, "MOIC", "moic")]:
    r0 = 4
    S.cell(row=r0, column=block_col, value=f"{label}: exit multiple (rows) vs. exit year").font = font(NAVY, True)
    header(S, r0 + 1, ["Exit EV / EBITDA"] + [y for y, _, _ in exit_years], block_col)
    for j, d in enumerate(deltas):
        rr = r0 + 2 + j
        mc = S.cell(row=rr, column=block_col, value=f"=LBO!$B${ENTRY_MULT}+({d})")
        mc.number_format = MULT
        mc.font = font(GREEN, d == 0)
        for k, (_, col, n) in enumerate(exit_years):
            eq = f"({L(block_col)}{rr}*LBO!${col}${EBITDA}-LBO!${col}${NET_DEBT})"
            f = f"=({eq}/LBO!$B${EQUITY})^(1/{n})-1" if kind == "irr" else f"={eq}/LBO!$B${EQUITY}"
            c = S.cell(row=rr, column=block_col + 1 + k, value=f)
            c.number_format = PCT if kind == "irr" else '0.00"x"'
            c.font = font(GREEN, d == 0 and n == 5)
            if d == 0 and n == 5:
                c.fill = KEY_FILL
S["A13"] = "Reading the grid: the yellow cell matches the base case on the LBO tab. Moving up one row = " \
           "selling at a lower multiple than you paid -- the main risk in any LBO."
S["A13"].font = font("595959", italic=True, size=9)

# =============================================================================
# SCREEN RESULTS  (values from screen_and_rank.py)
# =============================================================================
R = wb.create_sheet("Screen Results")
setup(R, [6, 38, 9, 36, 13, 12, 12, 12, 10, 10])
title(R, "Take-private screen: top 25 of 107 candidates, ranked by base-case IRR",
      "Values (not formulas) produced by screen_and_rank.py from SEC EDGAR XBRL data. "
      "Stride's row uses FY2025 data from the screen; the LBO tab uses FY2026 actuals.")
ranked = pd.read_csv("top25.csv")
header(R, 4, ["Rank", "Company", "Ticker", "Industry (SIC)", "Mkt cap ($mm)", "EBITDA ($mm)",
              "EBITDA margin", "Entry EV/EBITDA", "MOIC", "IRR"])
for i, row in ranked.iterrows():
    rr = 5 + i
    vals = [i + 1, row["name"], row["ticker"], str(row["industry"]).title(), row["market_cap"] / 1e6,
            row["ebitda"] / 1e6, row["margin"], row["entry_multiple"], row["moic"], row["irr"]]
    fmts = [None, None, None, None, '$#,##0', '$#,##0', PCT, MULT, '0.00"x"', PCT]
    for j, (v, fm) in enumerate(zip(vals, fmts)):
        c = R.cell(row=rr, column=j + 1, value=v)
        c.font = font(BLUE if j >= 4 else BLACK, row["ticker"] == "LRN")
        if fm:
            c.number_format = fm
        if row["ticker"] == "LRN":
            c.fill = KEY_FILL

fr = 5 + len(ranked) + 2
R.cell(row=fr, column=2, value="Screening funnel").font = font(NAVY, True)
funnel = [("US companies with FY2025 XBRL financials", 6154),
          ("Pass loose pre-screen (EBITDA $25M-$2B, FCF > 0, margin > 8%)", 1034),
          ("With a ticker and listed", 940), ("Debt data complete (manual checks applied)", 896),
          ("Excl. financials, utilities, oil & gas, shipping, LPs", 707), ("Market cap >= $250M", 570),
          ("EBITDA $50M-$1B", 457), ("EBITDA margin >= 15%", 263), ("FCF / EBITDA >= 40%", 216),
          ("Net debt / EBITDA <= 2.0x", 159), ("Revenue growth -5% to +40%", 133),
          ("Profitable both years, no EBIT spike", 107)]
for k, (lab, n) in enumerate(funnel):
    R.cell(row=fr + 1 + k, column=2, value=lab).font = font()
    c = R.cell(row=fr + 1 + k, column=5, value=n)
    c.font = font(BLUE)
    c.number_format = "#,##0"

# Summary first
SUM = wb.create_sheet("Summary", 0)
setup(SUM, [42, 16, 4, 60])
title(SUM, "Project Pathfinder -- Take-private of Stride, Inc. (NYSE: LRN)",
      "Illustrative LBO built from public filings. Not investment advice.")
header(SUM, 4, ["Key output", "Value", "", "Note"])
items = [("Offer price per share", f"=LBO!B{OFFER}", USD2, "30% premium to $82.39 (Oct 17, 2026)"),
         ("Entry enterprise value ($mm)", f"=LBO!B{ENTRY_EV}", '$#,##0', "Net of ~$958mm cash + securities"),
         ("Entry EV / FY2026 EBITDA", f"=LBO!B{ENTRY_MULT}", MULT, "GAAP EBITDA; SBC not added back"),
         ("New debt ($mm)", f"=LBO!B{DEBT0}", '$#,##0', "Constrained by the 40% minimum equity rule"),
         ("New debt / EBITDA", f"=LBO!B{DEBT_MULT}", MULT, ""),
         ("Sponsor equity ($mm)", f"=LBO!B{EQUITY}", '$#,##0', ""),
         ("Leverage at exit (debt / EBITDA)", f"=LBO!G{LEV}", MULT, "Shows how much debt the cash flow repays"),
         ("MOIC (5-year hold)", f"=LBO!B{MOIC}", '0.00"x"', "Exit at entry multiple -- no multiple expansion"),
         ("IRR (5-year hold)", f"=LBO!B{IRR_ROW}", PCT, "Typical PE target: 20%+"),
         ("Sources = uses check", f"=LBO!C{SOURCES}", None, "")]
for k, (lab, f, fm, note) in enumerate(items):
    rr = 5 + k
    SUM.cell(row=rr, column=1, value=lab).font = font()
    c = SUM.cell(row=rr, column=2, value=f)
    c.font = font(GREEN, lab.startswith(("IRR", "MOIC")))
    if fm:
        c.number_format = fm
    SUM.cell(row=rr, column=4, value=note).font = font("595959", size=9, italic=True)
SUM["A17"] = "Color key:  blue = hard-coded input  ·  black = formula  ·  green = link to another tab  ·  yellow = key output / lever"
SUM["A17"].font = font("595959", size=9)
SUM["A18"] = "To test a scenario, change the blue cells on the Assumptions tab (e.g. premium, max leverage, exit multiple change)."
SUM["A18"].font = font("595959", size=9)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != F:
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, color=c.font.color, italic=c.font.italic)

wb.save("Stride_LBO_Model.xlsx")
print("saved")
